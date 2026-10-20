"""
Informe semanal en Word (.docx) a partir de datos/eventos.csv.
Lo usa monitor.py los viernes (o con /semanal). También se puede lanzar a mano:
    python informe_semanal.py            -> últimos 7 días
    python informe_semanal.py 14         -> últimos 14 días
"""
import csv
import sys
from collections import Counter, defaultdict
from datetime import datetime, timedelta, timezone
from pathlib import Path
from zoneinfo import ZoneInfo

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

BASE = Path(__file__).resolve().parent
CSV_EVENTOS = BASE / "datos" / "eventos.csv"
CARPETA = BASE / "datos" / "informes"

DIAS = ["lunes", "martes", "miércoles", "jueves", "viernes", "sábado", "domingo"]
MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto",
         "septiembre", "octubre", "noviembre", "diciembre"]
NOMBRE_AMBITO = {"local": "Santander y Cantabria", "nacional": "Nacional con efecto en Santander",
                 "cuenta": "Cuentas de organizaciones"}
ORDEN_PRIO = {"alta": 3, "media": 2, "baja": 1}
MAX_POR_TEMA = {"otros": 25, "eventos": 25}     # el resto de temas: todos


def fecha_larga(d):
    return f"{DIAS[d.weekday()]} {d.day} de {MESES[d.month - 1]}"


def leer(desde):
    if not CSV_EVENTOS.exists():
        return []
    filas = []
    with CSV_EVENTOS.open(encoding="utf-8") as f:
        for r in csv.DictReader(f):
            try:
                det = datetime.fromisoformat(r["detectado_utc"])
            except (KeyError, ValueError):
                continue
            if det >= desde:
                r["_det"] = det
                r["_temas"] = [t for t in r.get("temas", "").split("|") if t] or ["otros"]
                filas.append(r)
    return filas


def enlace(parrafo, texto, url):
    """Añade un hipervínculo clicable al párrafo."""
    r_id = parrafo.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                                  is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F5FAD")
    rpr.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rpr.append(u)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = texto
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    h.append(run)
    parrafo._p.append(h)


def tabla(doc, cabecera, filas):
    t = doc.add_table(rows=1, cols=len(cabecera))
    t.style = "Light Grid Accent 1"
    for i, c in enumerate(cabecera):
        t.rows[0].cells[i].text = c
    for f in filas:
        celdas = t.add_row().cells
        for i, v in enumerate(f):
            celdas[i].text = str(v)
    doc.add_paragraph()


def generar(dias=7, tz_nombre="Europe/Madrid", temas_orden=None):
    tz = ZoneInfo(tz_nombre)
    fin = datetime.now(timezone.utc)
    ini = fin - timedelta(days=dias)
    filas = leer(ini)
    ini_l, fin_l = ini.astimezone(tz), fin.astimezone(tz)

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)

    doc.add_heading("Informe semanal de eventos — Santander y Cantabria", 0)
    p = doc.add_paragraph(f"Periodo: del {fecha_larga(ini_l)} al {fecha_larga(fin_l)} de {fin_l.year}")
    p.runs[0].italic = True
    doc.add_paragraph(f"Generado automáticamente el {fin_l.strftime('%d/%m/%Y a las %H:%M')} "
                      "a partir de prensa, GDELT y redes sociales públicas.").runs[0].font.size = Pt(9)

    if not filas:
        doc.add_paragraph("No hay eventos registrados en este periodo.")
    else:
        # ---------- resumen en cifras
        doc.add_heading("1. Resumen en cifras", 1)
        alta = [r for r in filas if r.get("prioridad") == "alta"]
        cont_tema = Counter(t for r in filas for t in r["_temas"])
        cont_amb = Counter(r.get("ambito", "") for r in filas)
        cont_red = Counter(r.get("red", "") for r in filas)
        cont_dia = Counter(r["_det"].astimezone(tz).date() for r in filas)
        doc.add_paragraph(f"Se registraron {len(filas)} noticias y publicaciones, de las cuales "
                          f"{len(alta)} fueron de prioridad alta (avisos inmediatos).")
        orden = (temas_orden or []) + sorted(set(cont_tema) - set(temas_orden or []))
        tabla(doc, ["Tema", "Nº"], [(t.capitalize(), cont_tema[t]) for t in orden if cont_tema.get(t)])
        tabla(doc, ["Ámbito", "Nº"], [(NOMBRE_AMBITO.get(a, a), n) for a, n in cont_amb.most_common()])
        tabla(doc, ["Fuente / red", "Nº"], cont_red.most_common())
        tabla(doc, ["Día", "Nº"], [(fecha_larga(d), cont_dia[d]) for d in sorted(cont_dia)])

        # ---------- cronología de lo importante
        doc.add_heading("2. Cronología de lo más importante", 1)
        if not alta:
            doc.add_paragraph("No hubo eventos de prioridad alta esta semana.")
        dia_actual = None
        for r in sorted(alta, key=lambda x: x["_det"]):
            d = r["_det"].astimezone(tz)
            if d.date() != dia_actual:
                dia_actual = d.date()
                doc.add_heading(fecha_larga(d).capitalize(), 3)
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(f"{d.strftime('%H:%M')} · {', '.join(r['_temas']).upper()} · ").bold = True
            enlace(p, r.get("titulo", ""), r.get("url", ""))
            p.add_run(f"  ({r.get('fuente', '')}, {NOMBRE_AMBITO.get(r.get('ambito'), '')})").italic = True

        # ---------- detalle por tema
        doc.add_heading("3. Detalle por tema", 1)
        por_tema = defaultdict(list)
        for r in filas:
            por_tema[r["_temas"][0]].append(r)
        for t in orden:
            items = por_tema.get(t)
            if not items:
                continue
            doc.add_heading(f"{t.capitalize()} ({len(items)})", 2)
            items.sort(key=lambda x: (ORDEN_PRIO.get(x.get("prioridad"), 0), x["_det"]), reverse=True)
            tope = MAX_POR_TEMA.get(t, len(items))
            for r in items[:tope]:
                d = r["_det"].astimezone(tz)
                p = doc.add_paragraph(style="List Bullet")
                p.add_run(f"{d.strftime('%d/%m %H:%M')} · ")
                enlace(p, r.get("titulo", ""), r.get("url", ""))
                p.add_run(f"  ({r.get('fuente', '')} · {r.get('red', '')})").italic = True
            if len(items) > tope:
                doc.add_paragraph(f"… y {len(items) - tope} más (ver datos/eventos.csv).").runs[0].italic = True

    doc.add_heading("Metodología", 1)
    doc.add_paragraph(
        "Monitorización automática cada 10–20 minutos de fuentes públicas: Google News (consultas locales "
        "y nacionales), RSS de medios locales, GDELT, Reddit, Mastodon y, en su caso, Bluesky, canales "
        "públicos de Telegram e Instagram (API oficial Business Discovery de Meta, solo cuentas "
        "profesionales públicas). Las noticias se clasifican por palabras clave en temas y prioridades, se "
        "filtran por zona geográfica y se eliminan duplicados entre medios por similitud de titulares. "
        "La clasificación es automática y puede contener errores. Solo se recogen publicaciones públicas "
        "de organizaciones y medios; no se analizan perfiles personales.")

    CARPETA.mkdir(parents=True, exist_ok=True)
    ruta = CARPETA / f"informe_semanal_{fin_l.strftime('%Y-%m-%d')}.docx"
    doc.save(ruta)
    return ruta, len(filas), len([r for r in filas if r.get("prioridad") == "alta"])


if __name__ == "__main__":
    d = int(sys.argv[1]) if len(sys.argv) > 1 else 7
    ruta, n, a = generar(d)
    print(f"Informe: {ruta} ({n} eventos, {a} importantes)")
